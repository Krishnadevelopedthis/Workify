"""PDF conversions for Workify that are best done in Python.

One-shot:  python -I workify_tools.py <command> <args...>
  prints one JSON line: {"ok": true, ...} or {"error": "..."} (exit code 1).
Worker:    python -I workify_tools.py serve
  reads {"id", "cmd", "args"} JSON lines on stdin and answers each with
  {"id", "result"} on stdout, so heavy libraries are imported only once.
"""

import json
import os
import sys


# BaseException subclasses, so library code's `except Exception` never swallows them.
class Done(BaseException):
    def __init__(self, data):
        self.data = data


class ToolFailure(BaseException):
    def __init__(self, message):
        self.message = message


def out(**data):
    raise Done(data)


def fail(message):
    raise ToolFailure(message)


# A pixmap costs width * height * n bytes, and a few hundred bytes of PDF can
# declare a page or an image of billions of pixels, so every decode is budgeted.
MAX_RENDER_PIXELS = 50_000_000  # one rendered page: 150 MB as RGB
MAX_IMAGE_PIXELS = 100_000_000  # one embedded image: a 600 dpi A3 scan is ~70 MP


def _key_int(doc, xref, key):
    kind, value = doc.xref_get_key(xref, key)
    if kind == "xref":  # e.g. /Width 12 0 R
        value = doc.xref_object(int(value.split()[0]), compressed=True)
    try:
        return int(float(value))
    except ValueError:
        return 0


def _jpx_size(doc, xref):
    """JPX images are decoded at the size in their codestream, not /Width /Height."""
    import struct

    data = doc.xref_stream_raw(xref) or b""
    at = data.find(b"\xff\x4f\xff\x51")  # SOC + SIZ marker
    if at < 0 or len(data) < at + 24:
        return 0
    x, y, x0, y0 = struct.unpack(">4I", data[at + 8 : at + 24])
    return max(x - x0, 0) * max(y - y0, 0)


def _image_pixels(doc, xref):
    """Pixels MuPDF would allocate to decode image `xref` (or its soft mask)."""
    worst = 0
    kind, smask = doc.xref_get_key(xref, "SMask")
    for x in (xref, int(smask.split()[0]) if kind == "xref" else 0):
        if x <= 0 or x >= doc.xref_length():
            continue
        px = _key_int(doc, x, "Width") * _key_int(doc, x, "Height")
        if "JPXDecode" in doc.xref_get_key(x, "Filter")[1]:
            px = max(px, _jpx_size(doc, x))
        worst = max(worst, px)
    return worst


def _check_page_images(doc, page):
    # Rendering or listing a page decodes its images, so refuse oversized ones up front.
    for img in page.get_images(full=True):
        if _image_pixels(doc, img[0]) > MAX_IMAGE_PIXELS:
            fail(f"Page {page.number + 1} contains an image that is too large to process.")


def open_pdf(path):
    import pymupdf

    try:
        doc = pymupdf.open(path)
    except Exception:
        fail("This file is not a valid PDF or is too damaged to read.")
    if doc.needs_pass:
        fail("This PDF is password protected. Unlock it first with Unlock PDF.")
    return doc


def pdf_to_docx(src, dst):
    from pdf2docx import Converter

    open_pdf(src).close()
    cv = Converter(src)
    try:
        cv.convert(dst)
    finally:
        cv.close()
    out(ok=True)


def pdf_to_xlsx(src, dst):
    import pdfplumber
    from openpyxl import Workbook

    open_pdf(src).close()
    wb = Workbook()
    wb.remove(wb.active)
    with pdfplumber.open(src) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            ws = wb.create_sheet(f"Page {i}")
            tables = page.extract_tables()
            row = 1
            if tables:
                for table in tables:
                    for cells in table:
                        for col, value in enumerate(cells, start=1):
                            ws.cell(row=row, column=col, value=_number(value))
                        row += 1
                    row += 1  # blank row between tables
            else:
                # No ruled table found: keep the text so no data is lost.
                for line in (page.extract_text() or "").splitlines():
                    for col, value in enumerate(line.split("  "), start=1):
                        if value.strip():
                            ws.cell(row=row, column=col, value=_number(value.strip()))
                    row += 1
    wb.save(dst)
    out(ok=True)


def _number(value):
    if value is None:
        return None
    text = str(value).strip()
    try:
        cleaned = text.replace(",", "")
        return int(cleaned) if cleaned.lstrip("-").isdigit() else float(cleaned)
    except ValueError:
        return text


def pdf_to_pptx(src, dst, scanned=frozenset()):
    """`scanned` holds 0-based pages whose text came from OCR: their full-page
    scan image is left out, so the slide shows the editable text instead."""
    import io

    import pymupdf
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Pt

    doc = open_pdf(src)
    prs = Presentation()
    first = doc[0].rect
    prs.slide_width = Emu(int(first.width * 12700))
    prs.slide_height = Emu(int(first.height * 12700))
    blank = prs.slide_layouts[6]
    pt = lambda v: Emu(int(v * 12700))  # PDF points -> EMU

    for pno, page in enumerate(doc):
        _check_page_images(doc, page)
        slide = prs.slides.add_slide(blank)
        page_area = page.rect.width * page.rect.height
        # Images first so text boxes sit on top of them.
        for info in page.get_image_info(xrefs=True):
            xref = info.get("xref")
            if not xref:
                continue
            x0, y0, x1, y1 = info["bbox"]
            if pno in scanned and (x1 - x0) * (y1 - y0) > 0.6 * page_area:
                continue
            try:
                pix = pymupdf.Pixmap(doc, xref)
                if pix.n - pix.alpha >= 4:
                    pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
                slide.shapes.add_picture(io.BytesIO(pix.tobytes("png")), pt(x0), pt(y0), pt(x1 - x0), pt(y1 - y0))
            except Exception:
                continue
        for block in page.get_text("dict", flags=pymupdf.TEXTFLAGS_DICT & ~pymupdf.TEXT_PRESERVE_IMAGES)["blocks"]:
            if block.get("type") != 0:
                continue
            x0, y0, x1, y1 = block["bbox"]
            box = slide.shapes.add_textbox(pt(x0), pt(y0), pt(max(x1 - x0, 10)), pt(max(y1 - y0, 10)))
            tf = box.text_frame
            tf.word_wrap = True
            for li, line in enumerate(block["lines"]):
                para = tf.paragraphs[0] if li == 0 else tf.add_paragraph()
                for span in line["spans"]:
                    run = para.add_run()
                    run.text = span["text"]
                    run.font.size = Pt(max(span["size"], 1))
                    run.font.bold = bool(span["flags"] & 16)
                    run.font.italic = bool(span["flags"] & 2)
                    c = 0 if pno in scanned else span["color"]  # OCR text is invisible in the PDF
                    run.font.color.rgb = RGBColor((c >> 16) & 255, (c >> 8) & 255, c & 255)
    prs.save(dst)
    out(ok=True)


def _ocr_layer(src, dst, lang):
    """Adds an OCR text layer to the pages that have no text. Returns the PDF to
    convert and the 0-based pages that were recognized (empty: nothing to do)."""
    import ocrmypdf

    doc = open_pdf(src)
    scanned = frozenset(i for i, page in enumerate(doc) if not page.get_text().strip())
    doc.close()
    if not scanned:
        return src, scanned
    layered = dst + ".ocr.pdf"
    try:
        ocrmypdf.ocr(src, layered, language=lang, skip_text=True, output_type="pdf", progress_bar=False, optimize=0)
    except ocrmypdf.exceptions.MissingDependencyError as e:
        fail(f"OCR engine missing: {e}")
    return layered, scanned


def _text_docx(src, dst):
    """A Word document built from the text blocks of each page (one page per section)."""
    import pymupdf
    from docx import Document
    from docx.shared import Pt

    doc = pymupdf.open(src)
    word = Document()
    for pno, page in enumerate(doc):
        if pno:
            word.add_page_break()
        for block in page.get_text("dict", flags=pymupdf.TEXTFLAGS_DICT & ~pymupdf.TEXT_PRESERVE_IMAGES)["blocks"]:
            if block.get("type") != 0:
                continue
            lines = [" ".join(s["text"] for s in line["spans"]).strip() for line in block["lines"]]
            text = "\n".join(line for line in lines if line)
            if not text:
                continue
            size = max((s["size"] for line in block["lines"] for s in line["spans"]), default=11)
            word.add_paragraph().add_run(text).font.size = Pt(min(max(round(size), 8), 36))
    word.save(dst)


def ocr_convert(kind, src, dst, lang):
    """PDF -> docx/pptx/xlsx for scanned documents: OCR first, then convert."""
    path, scanned = _ocr_layer(src, dst, lang)
    try:
        if kind == "docx":
            if scanned:
                _text_docx(path, dst)
            else:
                pdf_to_docx(src, dst)
        elif kind == "pptx":
            pdf_to_pptx(path, dst, scanned)
        elif kind == "xlsx":
            pdf_to_xlsx(path, dst)
        else:
            fail(f"Unknown output type: {kind}")
    finally:
        if path != src:
            try:
                os.remove(path)
            except OSError:
                pass  # still open on Windows; the job folder is deleted later anyway
    out(ok=True, recognized=len(scanned))


def pdf_to_md(src, dst):
    import pymupdf4llm

    open_pdf(src).close()
    md = pymupdf4llm.to_markdown(src)
    with open(dst, "w", encoding="utf-8") as f:
        f.write(md)
    out(ok=True)


def pdf_to_jpg(src, out_dir, dpi, stem):
    import pymupdf

    doc = open_pdf(src)
    files = []
    for i, page in enumerate(doc, start=1):
        _check_page_images(doc, page)
        zoom = int(dpi) / 72
        px = page.rect.width * zoom * page.rect.height * zoom
        if px > MAX_RENDER_PIXELS:
            zoom *= (MAX_RENDER_PIXELS / px) ** 0.5  # huge page: lower the dpi to fit
        pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom))
        name = f"{stem}-{i:03d}.jpg"
        pix.save(os.path.join(out_dir, name), jpg_quality=90)
        files.append(name)
    out(ok=True, files=files)


def extract_images(src, out_dir, stem):
    import pymupdf

    doc = open_pdf(src)
    files, seen = [], set()
    for pno, page in enumerate(doc, start=1):
        for img in page.get_images(full=True):
            xref = img[0]
            if xref in seen:
                continue
            seen.add(xref)
            if _image_pixels(doc, xref) > MAX_IMAGE_PIXELS:
                continue  # would decode to gigabytes
            pix = pymupdf.Pixmap(doc, xref)
            if pix.n - pix.alpha >= 4 or pix.alpha:
                pix = pymupdf.Pixmap(pymupdf.csRGB, pix) if pix.n - pix.alpha >= 4 else pymupdf.Pixmap(pix, 0)
            if pix.width < 16 or pix.height < 16:
                continue  # skip tiny decorations
            name = f"{stem}-p{pno}-{len(files) + 1:03d}.jpg"
            pix.save(os.path.join(out_dir, name), jpg_quality=92)
            files.append(name)
    if not files:
        fail("No images were found inside this PDF.")
    out(ok=True, files=files)


def redact(src, dst, spec_path):
    """spec: {"areas": [{"page": 1, "rect": [x0,y0,x1,y1]}], "terms": ["..."]}
    Rects are in PDF user space (bottom-left origin), as the browser sends them."""
    import pymupdf

    with open(spec_path, encoding="utf-8") as f:
        spec = json.load(f)
    doc = open_pdf(src)
    count = 0
    for area in spec.get("areas", []):
        idx = int(area["page"]) - 1
        if 0 <= idx < len(doc):
            page = doc[idx]
            rect = (pymupdf.Rect(*area["rect"]) * page.transformation_matrix).normalize()
            page.add_redact_annot(rect, fill=(0, 0, 0))
            count += 1
    for term in [t for t in spec.get("terms", []) if t.strip()]:
        for page in doc:
            for rect in page.search_for(term):
                page.add_redact_annot(rect, fill=(0, 0, 0))
                count += 1
    if count == 0:
        fail("Nothing to redact. Draw boxes or enter text to search for.")
    for page in doc:
        # Removes the underlying text, image pixels and vector graphics.
        page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_PIXELS, graphics=pymupdf.PDF_REDACT_LINE_ART_REMOVE_IF_TOUCHED)
    doc.set_metadata({})
    doc.save(dst, garbage=4, deflate=True)
    out(ok=True, redactions=count)


def ocr(src, dst, lang):
    import ocrmypdf

    open_pdf(src).close()
    try:
        ocrmypdf.ocr(src, dst, language=lang, skip_text=True, output_type="pdf", progress_bar=False, optimize=0)
    except ocrmypdf.exceptions.PriorOcrFoundError:
        fail("This PDF already contains text.")
    except ocrmypdf.exceptions.MissingDependencyError as e:
        fail(f"OCR engine missing: {e}")
    out(ok=True)


def repair(src, dst):
    import pymupdf

    try:
        doc = pymupdf.open(src)  # MuPDF rebuilds broken xref tables on open
    except Exception:
        fail("This PDF is too damaged to recover.")
    if doc.page_count == 0:
        fail("No pages could be recovered from this PDF.")
    doc.save(dst, garbage=4, deflate=True, clean=True)
    out(ok=True, pages=doc.page_count)


def detect_fields(src):
    """Finds likely form fields in a non-fillable PDF: underscore runs and
    label lines ending with a colon. Returns rects in PDF user space."""
    import re

    import pymupdf

    doc = open_pdf(src)
    fields = []
    for pno, page in enumerate(doc, start=1):
        to_pdf = ~page.transformation_matrix  # MuPDF space -> PDF user space
        found = []
        for x0, y0, x1, y1, text, *_ in page.get_text("words"):
            if re.fullmatch(r"_{4,}", text):
                found.append(([x0, y0 - 6, x1, y1], "text"))
            elif re.fullmatch(r"\[\s?\]|☐|□", text):
                found.append(([x0, y0, x1, y1], "checkbox"))
        for b in page.get_text("blocks"):
            line = b[4].strip()
            if line.endswith(":") and len(line) < 40 and "\n" not in line:
                x1, y0, y1 = b[2], b[1], b[3]
                width = min(200, page.rect.width - x1 - 36)
                if width > 40:
                    found.append(([x1 + 6, y0 - 2, x1 + 6 + width, y1 + 2], "text"))
        for rect, kind in found:
            r = (pymupdf.Rect(*rect) * to_pdf).normalize()
            fields.append({"page": pno, "rect": [r.x0, r.y0, r.x1, r.y1], "kind": kind})
    out(ok=True, fields=fields)


COMMANDS = {
    "pdf2docx": (pdf_to_docx, 2),
    "pdf2xlsx": (pdf_to_xlsx, 2),
    "pdf2pptx": (pdf_to_pptx, 2),
    "pdf2md": (pdf_to_md, 2),
    "ocr-convert": (ocr_convert, 4),
    "pdf2jpg": (pdf_to_jpg, 4),
    "extract-images": (extract_images, 3),
    "redact": (redact, 3),
    "ocr": (ocr, 3),
    "repair": (repair, 2),
    "detect-fields": (detect_fields, 1),
}

def run(cmd, args):
    """Runs one command and always returns a JSON-able dict."""
    if cmd not in COMMANDS:
        return {"error": f"Unknown command. Use one of: {', '.join(COMMANDS)}"}
    fn, arity = COMMANDS[cmd]
    if len(args) != arity or not all(isinstance(a, str) for a in args):
        return {"error": f"{cmd} expects {arity} text arguments"}
    try:
        fn(*args)
        return {"ok": True}
    except Done as d:
        return d.data
    except ToolFailure as f:
        return {"error": f.message}
    except ImportError as e:
        print(f"missing python package: {e}", file=sys.stderr)
        return {"error": "This conversion engine is not installed on the server yet."}
    except Exception as e:  # report every failure as JSON for the server
        return {"error": f"{type(e).__name__}: {e}"}


def limit_memory():
    """Backstop for decodes the checks above cannot see (third-party converters,
    images in patterns or annotations): past this much extra address space an
    allocation raises MemoryError in this worker instead of waking the OOM killer."""
    try:
        import resource
    except ImportError:
        return  # Windows has no rlimits
    try:
        budget = int(os.environ.get("WORKIFY_PY_MEMORY_MB") or 1536) * 2**20
    except ValueError:
        budget = 1536 * 2**20
    if budget <= 0:
        return
    try:
        # Relative to what the imports already mapped, much of it reserved, not used.
        with open("/proc/self/statm") as f:
            base = int(f.read().split()[0]) * os.sysconf("SC_PAGE_SIZE")
    except (OSError, ValueError):
        base = 0
    soft, hard = resource.getrlimit(resource.RLIMIT_AS)
    limit = base + budget
    if hard != resource.RLIM_INFINITY:
        limit = min(limit, hard)
    if soft == resource.RLIM_INFINITY or limit < soft:
        resource.setrlimit(resource.RLIMIT_AS, (limit, hard))


def serve():
    # Keep a private handle on the real stdout for protocol replies, then point
    # fd 1 at stderr so library prints (even from C code) cannot corrupt it.
    proto = os.fdopen(os.dup(1), "w", encoding="utf-8", buffering=1)
    os.dup2(2, 1)
    sys.stdout = sys.stderr
    # Warm up the heavy imports once, before the first request arrives.
    for mod in ("pymupdf", "pdfplumber", "pdf2docx", "pymupdf4llm", "openpyxl", "pptx"):
        try:
            __import__(mod)
        except Exception:
            pass
    limit_memory()
    proto.write(json.dumps({"ready": True}) + "\n")
    for line in sys.stdin:
        try:
            req = json.loads(line)
        except ValueError:
            continue
        result = run(str(req.get("cmd")), req.get("args") or [])
        proto.write(json.dumps({"id": req.get("id"), "result": result}, ensure_ascii=False) + "\n")


if __name__ == "__main__":
    if len(sys.argv) >= 2 and sys.argv[1] == "serve":
        serve()
    else:
        result = run(sys.argv[1] if len(sys.argv) > 1 else "", sys.argv[2:])
        print(json.dumps(result, ensure_ascii=False))
        sys.exit(1 if "error" in result else 0)
